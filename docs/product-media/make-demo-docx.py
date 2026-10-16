"""Rebuild the fictional DOCX used in the mdoc product captures (python-docx)."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).parent
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(9)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor.from_string('243E37')
doc.add_heading('Services agreement', 0)
doc.add_paragraph('ALDER STUDIO  /  CONSULTING SERVICES', 'Subtitle')
doc.add_heading('Parties', 1)
doc.add_paragraph('This agreement is between Alder Studio Ltd (the Client) and Maya Bennett (the Consultant).')
doc.add_paragraph('Maya Bennett will advise Alder Studio Ltd on the design and delivery of its customer research programme.')
doc.add_heading('Scope of work', 1)
doc.add_paragraph('The Consultant will review the research plan, lead six interviews, and prepare a written report with practical recommendations.')
table = doc.add_table(rows=1, cols=3)
table.style = 'Light Shading Accent 1'
for cell, label in zip(table.rows[0].cells, ['Milestone', 'Deliverable', 'Due']):
    cell.text = label
for row in [('Discovery', 'Research plan', '12 November'), ('Research', 'Six interviews', '26 November'), ('Delivery', 'Final report', '10 December')]:
    for cell, value in zip(table.add_row().cells, row):
        cell.text = value
doc.add_heading('Review and approval', 1)
doc.add_paragraph('Alder Studio Ltd will review each deliverable within five working days. Maya Bennett will address one round of consolidated feedback.')
doc.add_heading('Confidentiality', 1)
doc.add_paragraph('Each party will keep the other party’s confidential information private and use it only to perform this agreement.')
doc.add_heading('Contact', 1)
doc.add_paragraph('Project updates should be sent to maya.bennett@example.com.')
footer = section.footer.paragraphs[0]
footer.text = 'FICTIONAL AGREEMENT  ·  PRODUCT DEMONSTRATION'
footer.runs[0].font.size = Pt(8)
doc.core_properties.author = 'Lawkitt'
doc.core_properties.title = 'Services agreement — fictional demonstration'
doc.save(root / 'Services agreement.docx')
